import streamlit as st
import anthropic
import google.generativeai as genai
import re
import os
import io
import tempfile
import subprocess
from youtube_transcript_api import YouTubeTranscriptApi
from youtube_transcript_api._errors import (
    TranscriptsDisabled,
    NoTranscriptFound,
    VideoUnavailable,
)
from dotenv import load_dotenv
from googleapiclient.discovery import build
import markdown
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Try to import whisper (optional dependency)
try:
    import whisper
    WHISPER_AVAILABLE = True
except ImportError:
    WHISPER_AVAILABLE = False

load_dotenv()

st.set_page_config(
    page_title="Video Summarizer",
    page_icon="🎬",
    layout="wide",
)

st.title("Video Summarizer")
st.markdown("Summarize videos from YouTube, Vimeo, Twitter, TikTok, Bilibili, and 1000+ platforms.")


def extract_video_id(url: str) -> str | None:
    """Extract YouTube video ID from various URL formats."""
    patterns = [
        r'(?:youtube\.com\/watch\?v=)([a-zA-Z0-9_-]{11})',
        r'(?:youtube\.com\/embed\/)([a-zA-Z0-9_-]{11})',
        r'(?:youtu\.be\/)([a-zA-Z0-9_-]{11})',
        r'(?:youtube\.com\/shorts\/)([a-zA-Z0-9_-]{11})',
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return None


def is_youtube_url(url: str) -> bool:
    """Check if URL is a YouTube video."""
    return bool(re.search(r'(youtube\.com|youtu\.be)', url))


def extract_channel_id(url: str) -> str | None:
    """Extract YouTube channel ID or username from URL."""
    patterns = [
        r'(?:youtube\.com\/channel\/)([a-zA-Z0-9_-]+)',
        r'(?:youtube\.com\/c\/)([a-zA-Z0-9_-]+)',
        r'(?:youtube\.com\/@)([a-zA-Z0-9_-]+)',
        r'(?:youtube\.com\/user\/)([a-zA-Z0-9_-]+)',
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return None


def get_channel_videos(channel_identifier: str, youtube_api_key: str, max_results: int = 50) -> list[dict]:
    """Fetch video list from a YouTube channel."""
    youtube = build('youtube', 'v3', developerKey=youtube_api_key)

    # First, try to get channel ID if we have a username/handle
    if not channel_identifier.startswith('UC'):
        # Search for the channel
        search_response = youtube.search().list(
            q=channel_identifier,
            type='channel',
            part='id,snippet',
            maxResults=1
        ).execute()

        if not search_response.get('items'):
            return []

        channel_id = search_response['items'][0]['id']['channelId']
    else:
        channel_id = channel_identifier

    # Get the uploads playlist ID
    channel_response = youtube.channels().list(
        id=channel_id,
        part='contentDetails,snippet'
    ).execute()

    if not channel_response.get('items'):
        return []

    channel_title = channel_response['items'][0]['snippet']['title']
    uploads_playlist_id = channel_response['items'][0]['contentDetails']['relatedPlaylists']['uploads']

    # Fetch videos from uploads playlist
    videos = []
    next_page_token = None

    while len(videos) < max_results:
        playlist_response = youtube.playlistItems().list(
            playlistId=uploads_playlist_id,
            part='snippet',
            maxResults=min(50, max_results - len(videos)),
            pageToken=next_page_token
        ).execute()

        for item in playlist_response.get('items', []):
            videos.append({
                'video_id': item['snippet']['resourceId']['videoId'],
                'title': item['snippet']['title'],
                'channel': channel_title,
                'published_at': item['snippet']['publishedAt'],
            })

        next_page_token = playlist_response.get('nextPageToken')
        if not next_page_token:
            break

    return videos


def transcribe_with_whisper(video_url: str, model_size: str = "base") -> tuple[str, str]:
    """Download audio from any video URL and transcribe with Whisper.

    Returns: (transcript_text, language_code)
    """
    if not WHISPER_AVAILABLE:
        raise RuntimeError("Whisper transcription is not available. Please use a YouTube video with captions enabled.")

    with tempfile.TemporaryDirectory() as tmpdir:
        audio_path = os.path.join(tmpdir, "audio.mp3")

        # Download audio using yt-dlp (supports 1000+ sites)
        cmd = [
            "yt-dlp",
            "-x",  # Extract audio
            "--audio-format", "mp3",
            "--audio-quality", "0",
            "-o", audio_path,
            video_url
        ]

        # Add node runtime if available (required for YouTube)
        node_available = False
        try:
            node_check = subprocess.run(["which", "node"], capture_output=True, check=True)
            if node_check.returncode == 0:
                cmd.insert(1, "--js-runtimes")
                cmd.insert(2, "node")
                node_available = True
        except:
            pass

        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode != 0:
            error_msg = result.stderr
            if "403" in error_msg or "JavaScript runtime" in error_msg:
                raise RuntimeError(
                    "Audio download failed. This video requires captions to be available. "
                    "Try a different video with captions enabled, or run the app locally for full Whisper support."
                )
            if "Failed to resolve" in error_msg or "No address associated" in error_msg:
                raise RuntimeError(
                    "Network error: Cannot download audio in this environment. "
                    "This video has no captions available. Please try a YouTube video with captions enabled, "
                    "or run the app locally for full Whisper transcription support."
                )
            raise RuntimeError(f"Failed to download audio: {error_msg}")

        # Find the actual audio file (yt-dlp may add extension)
        audio_files = [f for f in os.listdir(tmpdir) if f.endswith(('.mp3', '.m4a', '.webm', '.opus', '.wav'))]
        if not audio_files:
            raise RuntimeError("No audio file found after download")

        actual_audio_path = os.path.join(tmpdir, audio_files[0])

        # Load Whisper model and transcribe
        model = whisper.load_model(model_size)
        result = model.transcribe(actual_audio_path)

        transcript_text = result["text"]
        language = result.get("language", "en")

        return transcript_text, language


def get_youtube_transcript(video_id: str, preferred_languages: list[str] = None) -> tuple[str, str]:
    """Fetch transcript from YouTube video with language preference.

    Returns: (transcript_text, language_code)
    """
    if preferred_languages is None:
        preferred_languages = ['en', 'zh-Hans', 'zh-Hant', 'zh', 'zh-CN', 'zh-TW']

    ytt_api = YouTubeTranscriptApi()

    # Try to fetch with preferred languages directly
    try:
        transcript = ytt_api.fetch(video_id, languages=preferred_languages)
        # Determine which language was used
        lang_code = 'en'  # default
        sample = " ".join([entry.text for entry in transcript[:10]]) if transcript else ""
        if any('\u4e00' <= char <= '\u9fff' for char in sample):
            lang_code = 'zh'

        full_text = " ".join([entry.text for entry in transcript])
        return full_text, lang_code
    except Exception:
        pass

    # Fallback: try listing available transcripts
    try:
        transcript_list = ytt_api.list(video_id)

        transcript = None
        lang_code = None

        for lang in preferred_languages:
            try:
                transcript = transcript_list.find_manually_created_transcript([lang])
                lang_code = lang
                break
            except NoTranscriptFound:
                continue

        if transcript is None:
            for lang in preferred_languages:
                try:
                    transcript = transcript_list.find_generated_transcript([lang])
                    lang_code = lang
                    break
                except NoTranscriptFound:
                    continue

        if transcript is None:
            for t in transcript_list:
                transcript = t
                lang_code = t.language_code
                break

        if transcript is None:
            raise NoTranscriptFound(video_id, preferred_languages, transcript_list)

        fetched = transcript.fetch()
        full_text = " ".join([entry.text for entry in fetched])
        return full_text, lang_code
    except NoTranscriptFound:
        raise
    except Exception as e:
        raise NoTranscriptFound(video_id, preferred_languages, None) from e


def get_transcript(video_url: str) -> tuple[str, str]:
    """Get transcript from any video URL.

    First tries YouTube captions if it's a YouTube URL, then falls back to Whisper.
    Returns: (transcript_text, language_code)
    """
    # Try YouTube captions first if it's a YouTube URL
    if is_youtube_url(video_url):
        video_id = extract_video_id(video_url)
        if video_id:
            try:
                return get_youtube_transcript(video_id)
            except (NoTranscriptFound, TranscriptsDisabled):
                pass  # Fall through to Whisper

    # Use Whisper for all other cases
    if WHISPER_AVAILABLE:
        return transcribe_with_whisper(video_url)
    else:
        raise RuntimeError("No transcript available and Whisper is not installed")


def summarize_with_gemini(transcript: str, api_key: str, source_language: str = "en", video_title: str = "") -> str:
    """Send transcript to Google Gemini API for summarization."""
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel('gemini-1.5-flash')

    language_note = ""
    if source_language.startswith('zh'):
        language_note = "\n\nNote: This transcript is in Chinese. Please provide the summary report entirely in English."

    title_context = ""
    if video_title:
        title_context = f"\n\nVideo Title: {video_title}"

    prompt = f"""Please analyze the following video transcript and create a comprehensive summary report in English.{title_context}{language_note}

Structure your report with these sections:

## Overview
A brief 2-3 sentence summary of what the video is about.

## Key Points
The main ideas or arguments presented in the video as bullet points.

## Detailed Summary
A more detailed breakdown of the content, organized by topic or chronologically.

## Key Takeaways
The most important insights or action items from the video.

## Notable Quotes
Any memorable or significant quotes from the video (if applicable).

---

TRANSCRIPT:
{transcript}"""

    response = model.generate_content(prompt)
    return response.text


def summarize_with_claude(transcript: str, api_key: str, source_language: str = "en", video_title: str = "") -> str:
    """Send transcript to Claude API for summarization."""
    client = anthropic.Anthropic(
        api_key=api_key,
        base_url="https://api.anthropic.com"
    )

    language_note = ""
    if source_language.startswith('zh'):
        language_note = "\n\nNote: This transcript is in Chinese. Please provide the summary report entirely in English."

    title_context = ""
    if video_title:
        title_context = f"\n\nVideo Title: {video_title}"

    message = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=4096,
        messages=[
            {
                "role": "user",
                "content": f"""Please analyze the following video transcript and create a comprehensive summary report in English.{title_context}{language_note}

Structure your report with these sections:

## Overview
A brief 2-3 sentence summary of what the video is about.

## Key Points
The main ideas or arguments presented in the video as bullet points.

## Detailed Summary
A more detailed breakdown of the content, organized by topic or chronologically.

## Key Takeaways
The most important insights or action items from the video.

## Notable Quotes
Any memorable or significant quotes from the video (if applicable).

---

TRANSCRIPT:
{transcript}""",
            }
        ],
    )
    return message.content[0].text


def summarize(transcript: str, api_key: str, ai_provider: str, source_language: str = "en", video_title: str = "") -> str:
    """Summarize transcript using selected AI provider."""
    if ai_provider == "Google Gemini (Free)":
        return summarize_with_gemini(transcript, api_key, source_language, video_title)
    else:
        return summarize_with_claude(transcript, api_key, source_language, video_title)


def convert_to_html(markdown_text: str, title: str = "Video Summary") -> str:
    """Convert markdown to styled HTML."""
    html_content = markdown.markdown(markdown_text, extensions=['tables', 'fenced_code'])

    styled_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Oxygen, Ubuntu, sans-serif;
            max-width: 800px;
            margin: 0 auto;
            padding: 40px 20px;
            line-height: 1.6;
            color: #333;
        }}
        h1 {{ color: #1a1a1a; border-bottom: 2px solid #e0e0e0; padding-bottom: 10px; }}
        h2 {{ color: #2c3e50; margin-top: 30px; }}
        ul {{ padding-left: 20px; }}
        li {{ margin-bottom: 8px; }}
        blockquote {{
            border-left: 4px solid #3498db;
            padding-left: 20px;
            margin-left: 0;
            color: #555;
            font-style: italic;
        }}
        hr {{ border: none; border-top: 1px solid #e0e0e0; margin: 30px 0; }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    {html_content}
</body>
</html>"""
    return styled_html


def convert_to_pdf(markdown_text: str, title: str = "Video Summary") -> bytes:
    """Convert markdown to PDF using weasyprint."""
    from weasyprint import HTML

    html_content = convert_to_html(markdown_text, title)
    pdf_bytes = HTML(string=html_content).write_pdf()
    return pdf_bytes


def convert_to_docx(markdown_text: str, title: str = "Video Summary") -> bytes:
    """Convert markdown to Word document."""
    doc = Document()

    title_para = doc.add_heading(title, 0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = markdown_text.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('---'):
            doc.add_paragraph('_' * 50)
        else:
            clean_line = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            clean_line = re.sub(r'\*(.+?)\*', r'\1', clean_line)
            if clean_line:
                doc.add_paragraph(clean_line)

    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    doc_bytes.seek(0)
    return doc_bytes.getvalue()


# Sidebar settings
gemini_api_key = os.getenv("GEMINI_API_KEY", "") or os.getenv("GOOGLE_API_KEY", "")
anthropic_api_key = os.getenv("ANTHROPIC_API_KEY", "")
youtube_api_key = os.getenv("YOUTUBE_API_KEY", "")

with st.sidebar:
    st.header("Settings")

    ai_provider = st.selectbox(
        "AI Provider",
        ["Google Gemini (Free)", "Anthropic Claude"],
        help="Gemini has a free tier. Claude requires paid credits."
    )

    if ai_provider == "Google Gemini (Free)":
        user_api_key = st.text_input(
            "Google AI API Key",
            value=gemini_api_key,
            type="password",
            help="Get free key at https://aistudio.google.com/app/apikey",
        )
    else:
        user_api_key = st.text_input(
            "Anthropic API Key",
            value=anthropic_api_key,
            type="password",
            help="Get key at https://console.anthropic.com/",
        )

    user_youtube_api_key = st.text_input(
        "YouTube Data API Key (optional)",
        value=youtube_api_key,
        type="password",
        help="Only needed for batch channel processing.",
    )

    st.markdown("---")
    st.markdown("### Supported Platforms")
    st.markdown("""
    - YouTube, Vimeo, Dailymotion
    - Twitter/X, TikTok, Instagram
    - Bilibili, Douyin, Weibo
    - Facebook, Twitch
    - And 1000+ more!
    """)

    st.markdown("---")
    st.markdown("### How to use")
    st.markdown("""
    1. Get a free API key
    2. Paste any video URL
    3. Click 'Summarize'
    """)

    st.markdown("---")
    st.caption("⚠️ Cloud version works best with YouTube videos that have captions. For videos without captions, run locally.")

# Main content - tabs for single vs batch
tab1, tab2 = st.tabs(["Single Video", "Batch (YouTube Channel)"])

with tab1:
    st.subheader("Summarize Any Video")
    video_url = st.text_input(
        "Video URL",
        placeholder="Paste any video URL (YouTube, TikTok, Twitter, Bilibili, etc.)",
        key="single_url"
    )

    if st.button("Summarize Video", type="primary", use_container_width=True):
        if not user_api_key:
            st.error(f"Please enter your {'Google AI' if ai_provider == 'Google Gemini (Free)' else 'Anthropic'} API key in the sidebar.")
        elif not video_url:
            st.error("Please enter a video URL.")
        else:
            try:
                with st.status("Processing video...", expanded=True) as status:
                    st.write("Fetching transcript...")

                    # Try YouTube captions first, then Whisper
                    transcript_source = "Whisper transcription"
                    if is_youtube_url(video_url):
                        video_id = extract_video_id(video_url)
                        if video_id:
                            try:
                                transcript, lang_code = get_youtube_transcript(video_id)
                                transcript_source = "YouTube captions"
                            except (NoTranscriptFound, TranscriptsDisabled, Exception):
                                if WHISPER_AVAILABLE:
                                    st.write("No captions found. Using Whisper to transcribe audio...")
                                    st.write("(This may take a few minutes)")
                                    transcript, lang_code = transcribe_with_whisper(video_url)
                                else:
                                    raise RuntimeError("No captions available and Whisper not installed")
                        else:
                            if WHISPER_AVAILABLE:
                                st.write("Using Whisper to transcribe audio...")
                                transcript, lang_code = transcribe_with_whisper(video_url)
                            else:
                                raise RuntimeError("Whisper not installed")
                    else:
                        # Non-YouTube URL - use Whisper directly
                        if WHISPER_AVAILABLE:
                            st.write("Downloading and transcribing audio with Whisper...")
                            st.write("(This may take a few minutes)")
                            transcript, lang_code = transcribe_with_whisper(video_url)
                        else:
                            raise RuntimeError("Whisper not installed. Required for non-YouTube videos.")

                    lang_display = "Chinese" if lang_code and lang_code.startswith('zh') else "English"
                    st.write(f"Transcript ready via {transcript_source} ({len(transcript):,} chars, {lang_display})")
                    st.write(f"Generating summary with {ai_provider}...")

                    summary = summarize(transcript, user_api_key, ai_provider, lang_code)
                    status.update(label="Complete!", state="complete", expanded=False)

                st.markdown("---")
                st.markdown("## Summary Report")
                st.markdown(summary)

                # Export buttons
                st.markdown("### Download Report")
                col1, col2, col3, col4 = st.columns(4)

                with col1:
                    st.download_button(
                        label="📄 Markdown",
                        data=summary,
                        file_name="video_summary.md",
                        mime="text/markdown",
                    )

                with col2:
                    html_content = convert_to_html(summary)
                    st.download_button(
                        label="🌐 HTML",
                        data=html_content,
                        file_name="video_summary.html",
                        mime="text/html",
                    )

                with col3:
                    try:
                        pdf_content = convert_to_pdf(summary)
                        st.download_button(
                            label="📕 PDF",
                            data=pdf_content,
                            file_name="video_summary.pdf",
                            mime="application/pdf",
                        )
                    except Exception as e:
                        st.button("📕 PDF", disabled=True, help=f"PDF generation failed: {e}")

                with col4:
                    docx_content = convert_to_docx(summary)
                    st.download_button(
                        label="📝 Word/Docs",
                        data=docx_content,
                        file_name="video_summary.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )

            except Exception as e:
                st.error(f"An error occurred: {type(e).__name__}: {str(e)}")

with tab2:
    st.subheader("Process YouTube Channel")
    st.info("Batch processing only works with YouTube channels.")

    channel_url = st.text_input(
        "YouTube Channel URL",
        placeholder="https://www.youtube.com/@ChannelName",
        key="channel_url"
    )

    max_videos = st.slider("Number of videos to process", min_value=1, max_value=50, value=5)

    if st.button("Process Channel", type="primary", use_container_width=True):
        if not user_api_key:
            st.error("Please enter your API key in the sidebar.")
        elif not user_youtube_api_key:
            st.error("Please enter your YouTube Data API key for channel processing.")
        elif not channel_url:
            st.error("Please enter a YouTube channel URL.")
        else:
            channel_id = extract_channel_id(channel_url)

            if not channel_id:
                st.error("Could not extract channel ID. Please check the URL.")
            else:
                try:
                    with st.status("Processing channel...", expanded=True) as status:
                        st.write("Fetching video list...")
                        videos = get_channel_videos(channel_id, user_youtube_api_key, max_videos)

                        if not videos:
                            st.error("No videos found.")
                        else:
                            st.write(f"Found {len(videos)} videos. Processing...")

                            all_summaries = []
                            progress_bar = st.progress(0)

                            for i, video in enumerate(videos):
                                st.write(f"Processing: {video['title'][:50]}...")
                                try:
                                    video_url = f"https://www.youtube.com/watch?v={video['video_id']}"
                                    transcript, lang_code = get_transcript(video_url)
                                    summary = summarize(
                                        transcript,
                                        user_api_key,
                                        ai_provider,
                                        lang_code,
                                        video['title']
                                    )
                                    all_summaries.append({
                                        'title': video['title'],
                                        'video_id': video['video_id'],
                                        'summary': summary,
                                        'language': lang_code,
                                    })
                                except Exception as e:
                                    all_summaries.append({
                                        'title': video['title'],
                                        'video_id': video['video_id'],
                                        'summary': f"Error: {str(e)}",
                                        'language': 'unknown',
                                    })

                                progress_bar.progress((i + 1) / len(videos))

                            status.update(label="Complete!", state="complete", expanded=False)

                    st.markdown("---")
                    st.markdown("## Channel Summary Report")

                    combined_report = f"# Channel Video Summaries\n\nProcessed {len(all_summaries)} videos\n\n"

                    for item in all_summaries:
                        combined_report += f"---\n\n# {item['title']}\n\n"
                        combined_report += f"**Video URL:** https://www.youtube.com/watch?v={item['video_id']}\n\n"
                        combined_report += f"{item['summary']}\n\n"

                    for item in all_summaries:
                        with st.expander(f"📺 {item['title'][:60]}..."):
                            st.markdown(f"[Watch Video](https://www.youtube.com/watch?v={item['video_id']})")
                            st.markdown(item['summary'])

                    st.markdown("### Download Combined Report")
                    col1, col2, col3, col4 = st.columns(4)

                    with col1:
                        st.download_button(
                            label="📄 Markdown",
                            data=combined_report,
                            file_name="channel_summaries.md",
                            mime="text/markdown",
                            key="batch_md"
                        )

                    with col2:
                        html_content = convert_to_html(combined_report, "Channel Summary Report")
                        st.download_button(
                            label="🌐 HTML",
                            data=html_content,
                            file_name="channel_summaries.html",
                            mime="text/html",
                            key="batch_html"
                        )

                    with col3:
                        try:
                            pdf_content = convert_to_pdf(combined_report, "Channel Summary Report")
                            st.download_button(
                                label="📕 PDF",
                                data=pdf_content,
                                file_name="channel_summaries.pdf",
                                mime="application/pdf",
                                key="batch_pdf"
                            )
                        except Exception:
                            st.button("📕 PDF", disabled=True, key="batch_pdf_disabled")

                    with col4:
                        docx_content = convert_to_docx(combined_report, "Channel Summary Report")
                        st.download_button(
                            label="📝 Word/Docs",
                            data=docx_content,
                            file_name="channel_summaries.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="batch_docx"
                        )

                except Exception as e:
                    st.error(f"An error occurred: {str(e)}")

st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray;'>Video Summarizer | Supports 1000+ platforms | Powered by Gemini & Whisper</div>",
    unsafe_allow_html=True,
)
